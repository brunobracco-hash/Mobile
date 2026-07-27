"""O .docx gerado precisa obedecer ao que o conversor do Kindle espera."""

import docx
import pytest
from docx.oxml.ns import qn

from pdf2kindle.converter import Options, convert


@pytest.fixture(scope="module")
def word(converted):
    return docx.Document(converted.output_path)


def test_titulos_usam_os_estilos_nativos(word):
    estilos = [p.style.name for p in word.paragraphs]
    assert "Heading 1" in estilos, "sem Heading 1 o Kindle não monta a navegação"


def test_metadados_viram_titulo_e_autor_do_ebook(word):
    assert word.core_properties.title == "A Página e a Tela"
    assert word.core_properties.author == "Maria Andrade"


def test_campo_de_sumario_presente(word):
    instrucoes = [el.text or "" for el in word.element.body.iter(qn("w:instrText"))]
    assert any("TOC" in texto for texto in instrucoes)


def test_capitulo_comeca_em_pagina_nova(word):
    capitulos = [
        p for p in word.paragraphs if p.style.name == "Heading 1" and p.text.startswith("Capítulo 2")
    ]
    assert capitulos and capitulos[0].paragraph_format.page_break_before


def test_sem_tabulacao_nem_recuo_manual(word):
    for p in word.paragraphs:
        assert "\t" not in p.text
        assert not p.text.startswith("   ")


def test_sem_paragrafos_vazios(word):
    vazios = [p for p in word.paragraphs if not p.text.strip() and not p.runs]
    assert not vazios, "parágrafo vazio vira linha solta no Kindle"


def test_sem_quebras_de_linha_manuais(word):
    assert not list(word.element.body.iter(qn("w:br")))


def test_sem_tabelas_nem_colunas(word):
    assert not word.tables
    section = word.sections[0]
    assert not list(section._sectPr.iter(qn("w:cols"))) or all(
        col.get(qn("w:num"), "1") == "1" for col in section._sectPr.iter(qn("w:cols"))
    )


def test_imagem_inline_e_dentro_da_largura_do_texto(word):
    assert len(word.inline_shapes) == 1
    largura_util = word.sections[0].page_width - word.sections[0].left_margin - word.sections[0].right_margin
    assert word.inline_shapes[0].width <= largura_util


def test_recuo_vem_do_estilo_e_nao_do_texto(word):
    corpo = [p for p in word.paragraphs if p.style.name == "Corpo Kindle"]
    assert corpo
    assert corpo[0].style.paragraph_format.first_line_indent is not None


def test_primeiro_paragrafo_apos_titulo_sem_recuo(word):
    estilos = [p.style.name for p in word.paragraphs]
    i = estilos.index("Heading 1")
    seguintes = [s for s in estilos[i + 1 :] if s.startswith("Corpo")]
    assert seguintes[0] == "Corpo Kindle Inicial"


def test_titulos_em_preto(word):
    assert word.styles["Heading 1"].font.color.rgb == docx.shared.RGBColor(0, 0, 0)


def test_notas_no_fim_do_documento(word):
    textos = [p.text for p in word.paragraphs]
    assert "Notas" in textos
    assert textos.index("Notas") > len(textos) - 4


def test_opcoes_desligam_rosto_sumario_e_imagens(sample_pdf, tmp_path):
    saida = tmp_path / "simples.docx"
    convert(
        sample_pdf,
        str(saida),
        Options(toc=False, title_page=False, keep_images=False, page_break_chapters=False, justify=False),
    )
    doc = docx.Document(str(saida))
    assert not doc.inline_shapes
    assert "Sumário" not in [p.text for p in doc.paragraphs]
    assert doc.paragraphs[0].style.name != "Title"


def test_titulo_e_autor_podem_ser_forcados(sample_pdf, tmp_path):
    saida = tmp_path / "forcado.docx"
    convert(sample_pdf, str(saida), Options(title="Outro Título", author="Outro Autor"))
    doc = docx.Document(str(saida))
    assert doc.core_properties.title == "Outro Título"
    assert doc.core_properties.author == "Outro Autor"
