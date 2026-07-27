"""Escrita do .docx pensado para o Kindle.

Regras que guiaram este módulo (o que o conversor da Amazon faz bem e o que ele
estraga):

  * O texto precisa ser **refluível**: nada de caixas de texto, colunas,
    tabulações, posicionamento absoluto ou quebras de linha manuais.
  * A navegação do e-book (o "Ir para") é construída a partir dos estilos
    Título 1/2/3 — por isso usamos os estilos nativos, e não texto grande em
    negrito.
  * Recuos e espaçamentos precisam vir do estilo do parágrafo. Espaços, tabs e
    parágrafos vazios viram lixo visível no Kindle.
  * Imagens devem ser inline e limitadas à largura do texto; imagem flutuante
    some ou desloca o texto.
  * Fonte e tamanho são substituídos pelo leitor, mas servem para a leitura no
    Word antes do envio.
"""

from __future__ import annotations

import io
import re
from typing import List, Optional, Set

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from .model import Document, Element, Run

HEADING_SIZES = {1: 20, 2: 15, 3: 13}


def _set_style_lang(style, lang: str) -> None:
    rpr = style.element.get_or_add_rPr()
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def _enable_widow_control(style) -> None:
    ppr = style.element.get_or_add_pPr()
    for tag in ("w:widowControl", "w:suppressAutoHyphens"):
        el = ppr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            ppr.append(el)
        el.set(qn("w:val"), "1" if tag == "w:widowControl" else "0")


def _update_fields_on_open(document) -> None:
    """Faz o Word calcular o sumário na primeira abertura."""
    settings = document.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def _add_toc_field(paragraph, levels: str = "1-3") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Abra no Word e atualize o campo (F9) para preencher o sumário."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def _configure_styles(document, font: str, body_pt: float, line_spacing: float, justify: bool, lang: str) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(body_pt)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    _set_style_lang(normal, lang)
    _enable_widow_control(normal)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT

    body = styles.add_style("Corpo Kindle", 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH
    body.base_style = normal
    body.paragraph_format.first_line_indent = Cm(0.5)
    body.paragraph_format.space_after = Pt(0)
    body.quick_style = True

    first = styles.add_style("Corpo Kindle Inicial", 1)
    first.base_style = body
    first.paragraph_format.first_line_indent = Cm(0)
    first.quick_style = True

    note = styles.add_style("Nota Kindle", 1)
    note.base_style = normal
    note.font.size = Pt(max(8, body_pt - 2))
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    figure = styles.add_style("Figura Kindle", 1)
    figure.base_style = normal
    figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.first_line_indent = Cm(0)
    figure.paragraph_format.space_before = Pt(10)
    figure.paragraph_format.space_after = Pt(10)

    for level, size in HEADING_SIZES.items():
        style = styles[f"Heading {level}"]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)   # o azul padrão do Word vaza para o Kindle
        hrpr = style.element.get_or_add_rPr()
        hfonts = hrpr.find(qn("w:rFonts"))
        if hfonts is None:
            hfonts = OxmlElement("w:rFonts")
            hrpr.append(hfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            hfonts.set(qn(attr), font)
        hpf = style.paragraph_format
        hpf.space_before = Pt(18 if level == 1 else 14)
        hpf.space_after = Pt(8)
        hpf.keep_with_next = True
        hpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hpf.first_line_indent = Cm(0)
        _set_style_lang(style, lang)

    title_style = styles["Title"]
    title_style.font.name = font
    title_style.font.size = Pt(28)
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _front_matter_duplicates(doc: Document) -> Set[int]:
    """Título e autor que abrem o PDF não precisam repetir a página de rosto."""
    normalize = lambda text: re.sub(r"\W+", "", text).lower()  # noqa: E731
    targets = {normalize(doc.title), normalize(doc.author)} - {""}
    skip: Set[int] = set()
    for element in doc.elements[:6]:
        if element.page > 0:
            break
        if element.kind in ("heading", "paragraph") and normalize(element.text) in targets:
            skip.add(id(element))
    return skip


def _add_runs(paragraph, runs: List[Run]) -> None:
    for run in runs:
        text = run.text
        if not text:
            continue
        r = paragraph.add_run(text)
        r.bold = run.bold
        r.italic = run.italic


def _normalize_image(data: bytes, ext: str) -> Optional[tuple]:
    """Garante um formato que o Word aceita; devolve (bytes, extensão)."""
    ext = (ext or "").lower()
    if ext in ("png", "jpeg", "jpg", "gif", "bmp", "tiff", "tif"):
        return data, ext
    try:  # jpx, jbig2 e afins: reconverte via PyMuPDF
        import fitz

        pix = fitz.Pixmap(data)
        if pix.n - pix.alpha >= 4:  # CMYK
            pix = fitz.Pixmap(fitz.csRGB, pix)
        return pix.tobytes("png"), "png"
    except Exception:
        return None


def _add_image(document, element: Element, content_width_cm: float) -> bool:
    prepared = _normalize_image(element.image or b"", element.image_ext)
    if not prepared:
        return False
    data, _ext = prepared
    paragraph = document.add_paragraph(style="Figura Kindle")
    width_pt, height_pt = element.image_size
    width_cm = content_width_cm
    if width_pt > 0:
        natural_cm = width_pt / 72 * 2.54
        width_cm = min(content_width_cm, natural_cm)
    try:
        paragraph.add_run().add_picture(io.BytesIO(data), width=Cm(width_cm))
    except Exception:
        p = paragraph._element
        p.getparent().remove(p)
        return False
    return True


def write_docx(
    doc: Document,
    output_path: str,
    font: str = "Georgia",
    body_pt: float = 12.0,
    line_spacing: float = 1.15,
    justify: bool = True,
    toc: bool = True,
    title_page: bool = True,
    page_break_chapters: bool = True,
    lang: str = "pt-BR",
) -> None:
    document = docx.Document()

    section = document.sections[0]
    section.page_width = Cm(15.0)      # página estreita: o texto quebra como em um livro
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    content_width_cm = 15.0 - 3.0

    _configure_styles(document, font, body_pt, line_spacing, justify, lang)
    _update_fields_on_open(document)

    document.core_properties.title = doc.title or "Documento"
    if doc.author:
        document.core_properties.author = doc.author
    document.core_properties.language = lang
    document.core_properties.comments = "Convertido de PDF por pdf2kindle"

    if title_page:
        p = document.add_paragraph(doc.title or "Documento", style="Title")
        p.paragraph_format.space_before = Pt(60)
        if doc.author:
            sub = document.add_paragraph(doc.author)
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub.paragraph_format.space_before = Pt(12)
            sub.runs[0].italic = True

    if toc:
        heading = document.add_heading("Sumário", level=1)
        heading.paragraph_format.page_break_before = True
        toc_paragraph = document.add_paragraph()
        toc_paragraph.paragraph_format.first_line_indent = Cm(0)
        _add_toc_field(toc_paragraph)

    skip = _front_matter_duplicates(doc) if title_page else set()

    first_content = True
    previous_kind = ""
    for element in doc.elements:
        if id(element) in skip:
            continue
        if element.kind == "heading":
            level = max(1, min(3, element.level or 1))
            paragraph = document.add_heading("", level=level)
            _add_runs(paragraph, element.runs)
            if level == 1 and page_break_chapters and not first_content:
                paragraph.paragraph_format.page_break_before = True
            elif (toc or title_page) and first_content:
                paragraph.paragraph_format.page_break_before = True
        elif element.kind == "paragraph":
            style = "Corpo Kindle Inicial" if previous_kind in ("heading", "image", "") else "Corpo Kindle"
            paragraph = document.add_paragraph(style=style)
            _add_runs(paragraph, element.runs)
        elif element.kind == "list_item":
            style = "List Number" if element.level == 1 else "List Bullet"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(2)
            _add_runs(paragraph, element.runs)
        elif element.kind == "note":
            paragraph = document.add_paragraph(style="Nota Kindle")
            _add_runs(paragraph, element.runs)
        elif element.kind == "image":
            if not _add_image(document, element, content_width_cm):
                continue
        else:
            continue
        previous_kind = element.kind
        first_content = False

    if doc.notes:
        heading = document.add_heading("Notas", level=1)
        heading.paragraph_format.page_break_before = True
        for note in doc.notes:
            paragraph = document.add_paragraph(style="Nota Kindle")
            marker = paragraph.add_run(f"[p. {note.page + 1}] ")
            marker.bold = True
            _add_runs(paragraph, note.runs)

    document.save(output_path)
